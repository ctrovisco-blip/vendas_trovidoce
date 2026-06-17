"""
Gera o relatório de Café em Grão (Buondi / Christina / Sical) a partir do
ficheiro Excel exportado do SAGE.

Uso manual:
    python gerar_relatorio.py

O script procura automaticamente o ficheiro Excel mais recente na pasta
PASTA_FONTE e guarda o relatório em PASTA_DESTINO.
"""

import sys
import glob
import os
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── Configuração ────────────────────────────────────────────────────────────
PASTA_FONTE   = r"I:\O meu disco\Trovidoce\Mapas Mensais"
PASTA_DESTINO = r"I:\O meu disco\Claude"
NOME_SAIDA    = "Relatorio_Grao_por_Vendedor.xlsx"

FAMILIAS = ["Buondi Grão Kg", "Christina Grão Kg", "Sical Grão Kg"]
KG_POR_CAIXA = 6
# ────────────────────────────────────────────────────────────────────────────


def encontrar_ficheiro_fonte():
    """Devolve o ficheiro .xlsx mais recente na pasta fonte."""
    padrao = os.path.join(PASTA_FONTE, "*.xlsx")
    ficheiros = glob.glob(padrao)
    if not ficheiros:
        raise FileNotFoundError(f"Nenhum ficheiro .xlsx encontrado em:\n  {PASTA_FONTE}")
    return max(ficheiros, key=os.path.getmtime)


def gerar(ficheiro_fonte=None):
    if ficheiro_fonte is None:
        ficheiro_fonte = encontrar_ficheiro_fonte()

    print(f"[{datetime.now():%H:%M:%S}] A ler: {os.path.basename(ficheiro_fonte)}")

    df = pd.read_excel(ficheiro_fonte, sheet_name="COLAR AQUI", header=1)

    df_f = df[df["Familia"].isin(FAMILIAS)].copy()
    df_f["Dt_ Emissao"] = pd.to_datetime(df_f["Dt_ Emissao"])
    df_f["Ano"] = df_f["Dt_ Emissao"].dt.year
    df_f["Mes"] = df_f["Dt_ Emissao"].dt.to_period("M")

    qtd = (
        df_f.groupby(["N_ Vend_", "Vendedor", "N_ Clie_", "Cliente", "Familia", "Ano"])["Quantidade"]
        .sum().reset_index()
    )

    ano_atual = datetime.now().year
    mes_atual = datetime.now().month

    def divisor_media(ano):
        if ano < ano_atual:
            return 12
        else:
            return mes_atual

    rel = qtd.copy()
    rel["Divisor"] = rel["Ano"].apply(divisor_media)
    rel["Media Mensal"] = (rel["Quantidade"] / rel["Divisor"]).round(1)
    rel = rel.sort_values(["N_ Vend_", "N_ Clie_", "Familia", "Ano"])

    # ── Excel ────────────────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Grão por Vendedor"

    header_fill   = PatternFill("solid", fgColor="1F4E79")
    subtotal_fill = PatternFill("solid", fgColor="D6E4F0")
    alt_fill      = PatternFill("solid", fgColor="EBF3FB")
    white_bold    = Font(bold=True, color="FFFFFF")
    bold          = Font(bold=True)
    center        = Alignment(horizontal="center")

    headers = [
        "Vendedor", "Cód. Cliente", "Nome Cliente", "Família", "Ano",
        "Qtd Total (cx)", "Qtd Total (kg)", "Média Mensal (cx)", "Média Mensal (kg)",
    ]
    widths = [28, 14, 52, 20, 8, 15, 15, 18, 18]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = white_bold
        cell.fill = header_fill
        cell.alignment = center
        ws.column_dimensions[get_column_letter(col)].width = w

    row_num = 2
    for vendedor in rel.sort_values("N_ Vend_")["Vendedor"].unique():
        df_v = rel[rel["Vendedor"] == vendedor]
        shade = False
        prev_cliente = None

        for _, r in df_v.iterrows():
            if r["N_ Clie_"] != prev_cliente:
                shade = not shade
                prev_cliente = r["N_ Clie_"]

            fill = alt_fill if shade else PatternFill()
            qtd_cx = int(r["Quantidade"])
            med_cx = r["Media Mensal"]

            row_data = [
                vendedor, int(r["N_ Clie_"]), r["Cliente"], r["Familia"], int(r["Ano"]),
                qtd_cx, qtd_cx * KG_POR_CAIXA, med_cx, round(med_cx * KG_POR_CAIXA, 1),
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col, value=val)
                cell.fill = fill
                if col >= 5:
                    cell.alignment = center
            row_num += 1


    ws.freeze_panes = "A2"

    os.makedirs(PASTA_DESTINO, exist_ok=True)
    saida = os.path.join(PASTA_DESTINO, NOME_SAIDA)
    wb.save(saida)
    print(f"[{datetime.now():%H:%M:%S}] Relatório guardado em:\n  {saida}")
    return saida


if __name__ == "__main__":
    ficheiro = sys.argv[1] if len(sys.argv) > 1 else None
    gerar(ficheiro)
