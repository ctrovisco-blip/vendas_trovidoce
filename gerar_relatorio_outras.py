"""
Gera relatório Excel de vendas por cliente/família (todas as famílias excepto café em grão)
usando a coluna "Familia Completa" do mapa fonte.

Uso manual:
    python gerar_relatorio_outras.py

Script independente de gerar_relatorio.py
"""

import sys
import glob
import os
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── Configuração ─────────────────────────────────────────────────────────────
PASTA_FONTE   = r"I:\O meu disco\Trovidoce\Mapas Mensais"
PASTA_DESTINO = r"I:\O meu disco\Claude"
NOME_SAIDA    = "Relatorio_Outras_Familias.xlsx"

EXCLUIR = [
    "Buondi Grão Kg", "Buondi Decafe Grão Kg",
    "Sical Grão Kg",  "Sical Decafe Grão Kg",
    "Christina Grão Kg", "Christina Decafe Grão Kg",
]
KG_POR_CAIXA = 6
DIAS_MES     = 30.44
# ─────────────────────────────────────────────────────────────────────────────


def encontrar_ficheiro_fonte():
    nome_esperado = "COLAR AQUI para Analise Familias.xlsx"
    caminho = os.path.join(PASTA_FONTE, nome_esperado)
    if os.path.exists(caminho):
        return caminho
    padrao = os.path.join(PASTA_FONTE, "*COLAR AQUI*.xlsx")
    ficheiros = glob.glob(padrao)
    if ficheiros:
        return max(ficheiros, key=os.path.getmtime)
    raise FileNotFoundError(
        f"Ficheiro '{nome_esperado}' não encontrado em:\n  {PASTA_FONTE}"
    )


def gerar(ficheiro_fonte=None):
    if ficheiro_fonte is None:
        ficheiro_fonte = encontrar_ficheiro_fonte()

    print(f"[{datetime.now():%H:%M:%S}] A ler: {os.path.basename(ficheiro_fonte)}")

    df = pd.read_excel(ficheiro_fonte, sheet_name="COLAR AQUI", header=1)

    # Filtrar famílias excluídas
    padrao_excluir = "|".join(EXCLUIR)
    df_f = df[~df["Familia Completa"].str.contains(padrao_excluir, na=False)].copy()
    df_f = df_f[df_f["Familia Completa"].notna()]

    df_f["Dt_ Emissao"] = pd.to_datetime(df_f["Dt_ Emissao"])
    df_f["Ano"] = df_f["Dt_ Emissao"].dt.year

    # Nome mais recente por código de cliente
    nome_atual = (
        df_f.sort_values("Dt_ Emissao")
        .groupby("N_ Clie_")["Cliente"].last()
        .reset_index().rename(columns={"Cliente": "Nome"})
    )
    df_f = df_f.merge(nome_atual, on="N_ Clie_")
    df_f["Cliente"] = df_f["Nome"]

    # Extrair nível mais detalhado da Familia Completa
    df_f["Familia_Det"] = df_f["Familia Completa"].apply(
        lambda x: x.split("\\")[-1] if "\\" in str(x) else str(x)
    )

    chave = ["N_ Vend_", "Vendedor", "N_ Clie_", "Cliente", "Familia Completa", "Familia_Det", "Ano"]

    qtd  = df_f.groupby(chave)["Quantidade"].sum().reset_index()
    span = df_f.groupby(chave)["Dt_ Emissao"].agg(
        primeira_compra="min", ultima_compra="max"
    ).reset_index()
    span["N_Meses"] = span.apply(
        lambda r: max((r["ultima_compra"] - r["primeira_compra"]).days / DIAS_MES, 1.0),
        axis=1
    )

    rel = qtd.merge(span[chave + ["N_Meses"]], on=chave)
    rel["Media Mensal"] = (rel["Quantidade"] / rel["N_Meses"]).round(1)
    rel = rel.sort_values(["Cliente", "Familia Completa", "Ano"])

    # ── Excel ─────────────────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Outras Famílias"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    alt_fill    = PatternFill("solid", fgColor="EBF3FB")
    white_bold  = Font(bold=True, color="FFFFFF")
    center      = Alignment(horizontal="center")

    headers = [
        "Vendedor", "Cód. Cliente", "Nome Cliente",
        "Família", "Ano",
        "Qtd Total (cx)", "Qtd Total (kg)",
        "Média Mensal (cx)", "Média Mensal (kg)",
    ]
    widths = [28, 14, 52, 32, 8, 15, 15, 18, 18]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = white_bold
        cell.fill = header_fill
        cell.alignment = center
        ws.column_dimensions[get_column_letter(col)].width = w

    row_num = 2
    shade = False
    prev_cliente = None

    for _, r in rel.iterrows():
        if r["N_ Clie_"] != prev_cliente:
            shade = not shade
            prev_cliente = r["N_ Clie_"]

        fill = alt_fill if shade else PatternFill()
        qtd_cx = int(r["Quantidade"])
        med_cx = r["Media Mensal"]

        row_data = [
            r["Vendedor"], int(r["N_ Clie_"]), r["Cliente"],
            r["Familia_Det"], int(r["Ano"]),
            qtd_cx, qtd_cx * KG_POR_CAIXA,
            med_cx, round(med_cx * KG_POR_CAIXA, 1),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.fill = fill
            if col >= 5:
                cell.alignment = center
        row_num += 1

    ws.freeze_panes = "A2"
    os.makedirs(PASTA_DESTINO, exist_ok=True)
    saida = os.path.join(PASTA_DESTINO, NOME_SAIDA)
    wb.save(saida)
    print(f"[{datetime.now():%H:%M:%S}] Relatório guardado em:\n  {saida}")
    return saida


if __name__ == "__main__":
    ficheiro = sys.argv[1] if len(sys.argv) > 1 else None
    gerar(ficheiro)
